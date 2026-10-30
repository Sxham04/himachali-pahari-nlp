"""Shareable files for mentors (WhatsApp/email): verse image, Excel sample, report PDF.

Devanagari is rendered by headless Chrome (matplotlib can't shape Indic scripts properly).
Run: uv run --with markdown python src/make_share.py   (after src/data.py)
Writes to share/ (git-ignored): bible_verses.png, bible_sample.xlsx, one PDF per report in REPORTS
"""
import html
import re
import subprocess
from pathlib import Path

import markdown
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from data import LANGS, NAMES, ROOT, read_clean, words

OUT = ROOT / "share"
CHROME = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
SAMPLE_N = 50
REPORTS = ["project1-report", "stage1-report", "stage2-report"]  # docs/<name>.md -> share/<name>.pdf
IMAGE_REFS = ["JHN.1.1", "LUK.6.31", "JHN.15.12"]  # short, well-known verses
FONT = "'Nirmala UI', 'Mangal', 'Noto Sans Devanagari', sans-serif"


def chrome(*args: str) -> None:
    subprocess.run([str(CHROME), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--hide-scrollbars", *args], check=True, capture_output=True, timeout=120)


def pick_sample(df: pd.DataFrame, n: int = SAMPLE_N) -> pd.DataFrame:
    """n train verses of 8-20 Hindi words, spread evenly over the train books."""
    pool = df[(df.split == "train") & df.hin.map(lambda t: 8 <= len(words(t)) <= 20)]
    return pool.iloc[:: max(1, len(pool) // n)].head(n)


def verse_image(df: pd.DataFrame) -> None:
    rows = df.set_index("ref").loc[IMAGE_REFS]
    cards = []
    for ref, r in rows.iterrows():
        lines = "".join(f"<tr><th>{NAMES[c]}</th><td>{html.escape(r[c])}</td></tr>" for c in LANGS)
        cards.append(f"<div class='card'><div class='ref'>{ref.replace('.', ' ', 1).replace('.', ':')}</div><table>{lines}</table></div>")
    page = f"""<html><head><meta charset="utf-8"><style>
      body {{ margin:0; background:#f6f4ef; font-family:{FONT}; width:900px; }}
      .wrap {{ padding:28px 32px; }}
      h1 {{ font-size:24px; margin:0 0 4px; color:#1f2a44; }}
      p.sub {{ margin:0 0 18px; color:#555; font-size:15px; }}
      .card {{ background:#fff; border-radius:10px; padding:14px 18px; margin-bottom:14px; box-shadow:0 1px 3px #0002; }}
      .ref {{ font-size:13px; color:#888; margin-bottom:6px; letter-spacing:.5px; }}
      table {{ border-collapse:collapse; width:100%; }}
      th {{ text-align:left; width:110px; color:#3b6ea8; font-size:15px; padding:5px 10px 5px 0; vertical-align:top; }}
      td {{ font-size:19px; padding:5px 0; line-height:1.5; }}
      tr + tr td, tr + tr th {{ border-top:1px solid #eee; }}
      .foot {{ color:#777; font-size:13px; }}
    </style></head><body><div class="wrap" id="w">
      <h1>One verse, four languages</h1>
      <p class="sub">Verse-aligned New Testament: Hindi, Mandyali, Bilaspuri, Kangri · 7,346 clean verses</p>
      {''.join(cards)}
      <div class="foot">Public Bible translations · cleaned and aligned by Soham Sharma</div>
    </div></body></html>"""
    src = OUT / "_verses.html"
    src.write_text(page, encoding="utf-8")
    chrome(f"--screenshot={OUT / 'bible_verses.png'}", "--window-size=900,840", "--force-device-scale-factor=2", src.as_uri())
    src.unlink()


def excel_sample(df: pd.DataFrame) -> None:
    counts = df.split.value_counts()
    about = [
        ("Bible parallel dataset: Hindi, Mandyali, Bilaspuri, Kangri", None),
        ("", None),
        ("What", "The same New Testament verses in four languages, aligned verse by verse."),
        ("Languages", "Hindi (hin), Mandyali (mjl, ISO name Mandeali), Bilaspuri (kfs), Kangri (xnr)"),
        ("Size", f"{len(df):,} clean verses (from 7,567 raw); all 27 New Testament books"),
        ("Cleaning", "Removed 3 empty rows, 13 rows with other verses pasted into a cell, 1 near-empty row, "
                     "and 204 misaligned verses (mostly a shifted Bilaspuri block in 1 Corinthians 3-6 and 15-16). "
                     "Mandyali's Latin 'l' full stop replaced by '।'; Hindi cross-references removed."),
        ("Split (by book)", f"train {counts.get('train', 0):,} (25 books) · dev {counts.get('dev', 0):,} (Romans) · "
                            f"test {counts.get('test', 0):,} (John)"),
        ("Avg words/verse", " · ".join(f"{NAMES[c]} {df[c].map(lambda t: len(words(t))).mean():.1f}" for c in LANGS)),
        ("Licences", "Hindi: IRV 2019, Bridge Connectivity Solutions, CC BY-SA 4.0. Kangri: The Love Fellowship, "
                     "CC BY-SA 4.0. Mandyali, Bilaspuri: being confirmed. Shared here for research feedback only."),
        ("Note", "The Mandyali text is the public Bible translation."),
        ("This file", f"'Sample' sheet: {SAMPLE_N} verses from the training books, spread across books."),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    for r in about:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 95

    s = wb.create_sheet("Sample")
    s.append(["Verse", *[NAMES[c] for c in LANGS]])
    for r in pick_sample(df).itertuples():
        s.append([r.ref, *[getattr(r, c) for c in LANGS]])
    head = PatternFill("solid", fgColor="1F2A44")
    for c in s[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name="Nirmala UI", size=11)
    s.column_dimensions["A"].width = 11
    for col in "BCDE":
        s.column_dimensions[col].width = 45
    s.freeze_panes = "B2"
    wb.save(OUT / "bible_sample.xlsx")


LIST_ITEM = re.compile(r"\s*(\d+\.|[-*]) ")


def spaced_lists(md: str) -> str:
    """Python-Markdown needs a blank line before a list (GitHub does not): add one where missing."""
    out, prev = [], ""
    for line in md.split("\n"):
        if LIST_ITEM.match(line) and prev.strip() and not LIST_ITEM.match(prev) and not prev.lstrip().startswith("|"):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def report_pdf(name: str) -> None:
    md = (ROOT / "docs" / f"{name}.md").read_text(encoding="utf-8")
    md = spaced_lists(md)
    body = markdown.markdown(md, extensions=["tables", "sane_lists"])
    page = f"""<html><head><meta charset="utf-8"><style>
      @page {{ size: A4; margin: 16mm 15mm; }}
      body {{ font-family:{FONT}; font-size:11pt; line-height:1.5; color:#222; }}
      h1 {{ font-size:20pt; color:#1f2a44; margin-bottom:2px; }}
      h3 {{ page-break-after:avoid; }}
      h2 {{ font-size:14pt; color:#1f2a44; border-bottom:1px solid #ddd; padding-bottom:3px; margin-top:22px;
            page-break-after:avoid; }}
      table {{ border-collapse:collapse; margin:8px 0; font-size:10pt; page-break-inside:avoid; }}
      th, td {{ border:1px solid #ccc; padding:4px 8px; vertical-align:top; }}
      th {{ background:#eef2f8; }}
      img {{ max-width:78%; display:block; margin:8px auto; page-break-inside:avoid; }}
      code {{ background:#f3f3f3; padding:0 3px; border-radius:3px; }}
      hr {{ border:none; }}
      a {{ color:#3b6ea8; text-decoration:none; }}
    </style></head><body>{body}</body></html>"""
    src = ROOT / "docs" / "_report.html"  # in docs/ so ../results/ image paths resolve
    src.write_text(page, encoding="utf-8")
    chrome(f"--print-to-pdf={OUT / f'{name}.pdf'}", src.as_uri())
    src.unlink()


def main() -> None:
    OUT.mkdir(exist_ok=True)
    df = read_clean()
    verse_image(df)
    excel_sample(df)
    for name in REPORTS:
        report_pdf(name)
    for f in sorted(OUT.iterdir()):
        print(f.name, f.stat().st_size)


if __name__ == "__main__":
    toy = pd.DataFrame({"split": ["train"] * 3, "hin": ["क ख ग घ ङ च छ ज", "क", "क ख ग घ ङ च छ ज झ"]})
    assert len(pick_sample(toy, 5)) == 2
    main()
